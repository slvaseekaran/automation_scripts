import requests
from requests.auth import HTTPBasicAuth
from collections import defaultdict
from docx import Document
from docx.shared import Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE  # Import the RELATIONSHIP_TYPE constant

# Helper function to add hyperlink
def add_hyperlink(paragraph, url, text):
    """
    Add a hyperlink to a paragraph.

    :param paragraph: The paragraph where the hyperlink should be added.
    :param url: The URL for the hyperlink.
    :param text: The text that will be clickable.
    """
    part = paragraph.part
    r_id = part.relate_to(url, RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
    hyperlink = OxmlElement('w:hyperlink')
    hyperlink.set(qn('r:id'), r_id)

    new_run = OxmlElement('w:r')
    rPr = OxmlElement('w:rPr')
    new_run.append(rPr)

    new_run_text = OxmlElement('w:t')
    new_run_text.text = text
    new_run.append(new_run_text)

    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)

    return paragraph

# Set your personal access token, organization, and project name
pat = "Cfyuz84bwahhr3Wh4XdhHsWfELe8o451CL35BWiQTToVWJyD0XrmJQQJ99BAACAAAAA1kgrNAAASAZDOMJr3"  # Replace with your PAT
organization = "rapyuta-robotics"
project = "sootballs"
base_url = f"https://dev.azure.com/{organization}/{project}/_apis"

# Query ID for the specific query
query_id = "896d1f72-4e72-4751-8345-fe52b2f1608f"

# API endpoint to execute the query and get work item IDs
query_work_items_url = f"{base_url}/wit/wiql/{query_id}?api-version=7.0"

# Make the request to execute the query
response = requests.get(query_work_items_url, auth=HTTPBasicAuth('', pat))

# Create a new Word document
doc = Document()
doc.add_heading("Work Items by Assignee", level=1)

# Check if the request was successful
if response.status_code == 200:
    # Extract work item IDs from the response
    work_item_ids = [item['id'] for item in response.json().get('workItems', [])]
    
    # Fetch details of each work item
    if work_item_ids:
        ids_str = ",".join(map(str, work_item_ids))
        work_items_details_url = f"{base_url}/wit/workitems?ids={ids_str}&api-version=7.0"
        
        details_response = requests.get(work_items_details_url, auth=HTTPBasicAuth('', pat))
        
        if details_response.status_code == 200:
            work_items = details_response.json().get('value', [])
            # Dictionary to hold work items grouped by assignee
            assignee_dict = defaultdict(list)
            
            for work_item in work_items:
                assignee = work_item['fields'].get('System.AssignedTo', {}).get('displayName', 'Unassigned')
                title = work_item['fields']['System.Title']
                state = work_item['fields']['System.State']
                work_item_url = f"https://dev.azure.com/{organization}/{project}/_workitems/edit/{work_item['id']}"
                assignee_dict[assignee].append((title, state, work_item_url))
            
            # Write the output to the Word document
            for assignee, items in assignee_dict.items():
                doc.add_heading(assignee, level=2)
                for i, (title, state, url) in enumerate(items, 1):
                    p = doc.add_paragraph(f"{i}. ")
                    r = p.add_run(f"{title} ")
                    r.font.size = Pt(12)
                    
                    # Add hyperlink for the work item URL
                    add_hyperlink(p, url, "Link")
                    
                    p.add_run(f" ({state})").font.size = Pt(12)
                doc.add_paragraph("\n")
            
            # Save the document
            doc_path = "sre.docx"
            doc.save(doc_path)
            print(f"Document saved as {doc_path}")
        else:
            print(f"Failed to fetch work item details: {details_response.status_code} - {details_response.text}")
    else:
        print("No work items found in the query.")
else:
    print(f"Failed to execute query: {response.status_code} - {response.text}")
